import csv
import io
import json
import os
import shutil
import threading
import re
import unicodedata
import time
from datetime import date, datetime, timezone
from http.server import ThreadingHTTPServer, SimpleHTTPRequestHandler
from pathlib import Path
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parent
INDEX = ROOT / 'templates' / 'index.html'
INITIAL_DATA = ROOT / 'data' / 'initial_dataset.csv'
MAX_UPLOAD_BYTES = int(os.environ.get('MAX_UPLOAD_MB', '50')) * 1024 * 1024

FIELDNAMES = [
    'rm', 'reg', 'pedido', 'coleta', 'estcol', 'envio', 'tipo', 'oprec', 'horaop',
    'baseent', 'uf', 'motivo', 'marca', 'transf', 'interceptado',
    'indenizado', 'finalizado'
]


CANON_MAP = {
    'rm':'rm','regional':'reg','reg':'reg',
    'numerodepedidojms':'pedido','numerodopedidojms':'pedido','pedido':'pedido','numeropedido':'pedido',
    'tempodecoleta':'coleta','coletadoem':'coleta','coleta':'coleta',
    'estacaodecoleta':'estcol','estcol':'estcol',
    'horadeenvio':'envio',
    'tipodaultimaoperacao':'tipo','tipo':'tipo',
    'operacaomaisrecente':'oprec','oprec':'oprec',
    'horariodaultimaoperacao':'horaop','horaop':'horaop',
    'basedeentrega':'baseent','baseent':'baseent',
    'ufdestino':'uf','uf':'uf',
    'motivosdaanomalia':'motivo','motivodaanomalia':'motivo','motivo':'motivo',
    'marcadeassinatura':'marca','marca':'marca',
    'marcadetransferenciaedevolucao':'transf','transf':'transf',
    'interceptado':'interceptado','interceptada':'interceptado','interceptacao':'interceptado',
    'statusdeinterceptacao':'interceptado','statusinterceptacao':'interceptado',
    'indenizado':'indenizado','indenizada':'indenizado','indenizacao':'indenizado',
    'statusdeindenizacao':'indenizado','statusindenizacao':'indenizado',
    'finalizado':'finalizado','finalizada':'finalizado','finalizacao':'finalizado',
    'statusfinalizacao':'finalizado','statusdefinalizacao':'finalizado',
}


def norm_header(value) -> str:
    text = '' if value is None else str(value)
    text = unicodedata.normalize('NFD', text)
    text = ''.join(ch for ch in text if unicodedata.category(ch) != 'Mn')
    return re.sub(r'[^a-z0-9]', '', text.lower())


def _strict_dt(y, m, d, hh=0, mm=0, ss=0):
    try:
        y, m, d = int(y), int(m), int(d)
        hh, mm, ss = int(hh or 0), int(mm or 0), int(ss or 0)
        if y < 100:
            y += 1900 if y >= 70 else 2000
        return datetime(y, m, d, hh, mm, ss)
    except Exception:
        return None


def _detect_date_order(values):
    dmy = mdy = 0
    for value in values:
        if value is None or isinstance(value, (datetime, date, int, float)):
            continue
        m = re.search(r'(?<!\d)(\d{1,2})\s*[-/.]\s*(\d{1,2})\s*[-/.]\s*\d{2,4}', str(value))
        if not m:
            continue
        a, b = int(m.group(1)), int(m.group(2))
        if a > 12 >= b:
            dmy += 1
        elif b > 12 >= a:
            mdy += 1
    return 'MDY' if mdy > dmy else 'DMY'


def normalize_date_value(value, order='DMY', epoch=None):
    if value is None or value == '':
        return ''
    if isinstance(value, datetime):
        dt = value
    elif isinstance(value, date):
        dt = datetime(value.year, value.month, value.day)
    elif isinstance(value, (int, float)):
        try:
            from openpyxl.utils.datetime import from_excel
            dt = from_excel(value, epoch=epoch) if epoch is not None else from_excel(value)
            if isinstance(dt, date) and not isinstance(dt, datetime):
                dt = datetime(dt.year, dt.month, dt.day)
        except Exception:
            dt = None
    else:
        raw = str(value).strip().strip('"\'')
        if not raw:
            return ''
        # Extrai a data mesmo quando ela vem dentro de um texto/fórmula.
        iso = re.search(r'(\d{4})\s*[-/.]\s*(\d{1,2})\s*[-/.]\s*(\d{1,2})(?:[ T]+(\d{1,2})(?::(\d{1,2}))?(?::(\d{1,2}))?)?', raw)
        dt = None
        if iso:
            dt = _strict_dt(*iso.groups())
        if dt is None:
            amb = re.search(r'(\d{1,2})\s*[-/.]\s*(\d{1,2})\s*[-/.]\s*(\d{2,4})(?:[ T]+(\d{1,2})(?::(\d{1,2}))?(?::(\d{1,2}))?\s*(AM|PM)?)?', raw, re.I)
            if amb:
                a, b, y, hh, mm, ss, ap = amb.groups()
                a, b, hh = int(a), int(b), int(hh or 0)
                if ap:
                    ap = ap.upper()
                    if ap == 'PM' and hh < 12: hh += 12
                    if ap == 'AM' and hh == 12: hh = 0
                if a > 12 and b <= 12:
                    day, month = a, b
                elif b > 12 and a <= 12:
                    month, day = a, b
                elif order == 'MDY':
                    month, day = a, b
                else:
                    day, month = a, b
                dt = _strict_dt(y, month, day, hh, mm or 0, ss or 0)
        if dt is None and re.fullmatch(r'\d{4,6}(?:[.,]\d+)?', raw):
            try:
                from openpyxl.utils.datetime import from_excel
                n = float(raw.replace(',', '.'))
                if 20000 <= n <= 100000:
                    dt = from_excel(n, epoch=epoch) if epoch is not None else from_excel(n)
            except Exception:
                dt = None
    if not dt:
        return ''
    return dt.strftime('%Y-%m-%d %H:%M:%S')


def _cell_text(cell, cached_cell=None):
    values = []
    if cached_cell is not None and cached_cell.value not in (None, ''):
        values.append(cached_cell.value)
    if cell is not None and cell.value not in (None, ''):
        values.append(cell.value)
    for value in values:
        if isinstance(value, (datetime, date)):
            return value.strftime('%Y-%m-%d %H:%M:%S') if isinstance(value, datetime) else value.isoformat()
        return str(value).strip()
    return ''


def parse_xlsx_rows(raw_bytes: bytes):
    from openpyxl import load_workbook
    from io import BytesIO

    wb_formula = load_workbook(BytesIO(raw_bytes), read_only=True, data_only=False)
    wb_values = load_workbook(BytesIO(raw_bytes), read_only=True, data_only=True)
    ws_f = wb_formula.worksheets[0]
    ws_v = wb_values.worksheets[0]

    # Acha o cabeçalho pela linha física real do Excel.
    best_row, best_score = 1, -1
    for row_num in range(1, min(ws_f.max_row, 20) + 1):
        score = 0
        for col_num in range(1, min(ws_f.max_column, 80) + 1):
            text = _cell_text(ws_f.cell(row_num, col_num), ws_v.cell(row_num, col_num))
            if CANON_MAP.get(norm_header(text)):
                score += 1
        if score > best_score:
            best_row, best_score = row_num, score

    headers = {}
    horaop_col = None
    for col_num in range(1, ws_f.max_column + 1):
        text = _cell_text(ws_f.cell(best_row, col_num), ws_v.cell(best_row, col_num))
        canon = CANON_MAP.get(norm_header(text))
        if canon:
            headers[col_num] = canon
            if canon == 'horaop':
                horaop_col = col_num

    coleta_samples = []
    envio_samples = []
    horaop_samples = []
    for row_num in range(best_row + 1, ws_f.max_row + 1):
        coleta_samples.append(_cell_text(ws_f.cell(row_num, 4), ws_v.cell(row_num, 4)))
        envio_samples.append(_cell_text(ws_f.cell(row_num, 7), ws_v.cell(row_num, 7)))
        if horaop_col:
            horaop_samples.append(_cell_text(ws_f.cell(row_num, horaop_col), ws_v.cell(row_num, horaop_col)))
    coleta_order = _detect_date_order(coleta_samples)
    envio_order = _detect_date_order(envio_samples)
    horaop_order = _detect_date_order(horaop_samples)

    rows = []
    envio_nonempty = envio_parsed = 0
    envio_sample = []
    epoch = wb_values.epoch

    for row_num in range(best_row + 1, ws_f.max_row + 1):
        row = {key: '' for key in FIELDNAMES}
        for col_num, canon in headers.items():
            row[canon] = _cell_text(ws_f.cell(row_num, col_num), ws_v.cell(row_num, col_num))

        # REGRA FIXA DO RELATÓRIO: D = coleta, G = hora de envio.
        d_raw = _cell_text(ws_f.cell(row_num, 4), ws_v.cell(row_num, 4))
        g_raw = _cell_text(ws_f.cell(row_num, 7), ws_v.cell(row_num, 7))
        row['coleta'] = normalize_date_value(ws_v.cell(row_num, 4).value if ws_v.cell(row_num, 4).value not in (None, '') else d_raw, coleta_order, epoch)
        row['envio'] = normalize_date_value(ws_v.cell(row_num, 7).value if ws_v.cell(row_num, 7).value not in (None, '') else g_raw, envio_order, epoch)
        if horaop_col:
            h_raw = _cell_text(ws_f.cell(row_num, horaop_col), ws_v.cell(row_num, horaop_col))
            h_value = ws_v.cell(row_num, horaop_col).value
            row['horaop'] = normalize_date_value(h_value if h_value not in (None, '') else h_raw, horaop_order, epoch)

        if g_raw:
            envio_nonempty += 1
            if len(envio_sample) < 5:
                envio_sample.append(f'G{row_num}={g_raw}')
        if row['envio']:
            envio_parsed += 1

        pedido = (row.get('pedido') or '').strip()
        if pedido:
            row['pedido'] = pedido
            row['reg'] = (row.get('reg') or '').strip().upper()
            rows.append(row)

    if envio_nonempty and not envio_parsed:
        raise ValueError(
            f'Coluna G contém {envio_nonempty} valores, mas nenhum horário foi interpretado. '
            f'Amostra: {"; ".join(envio_sample) or "sem amostra"}'
        )
    return rows, {
        'headerRow': best_row,
        'envioNonEmpty': envio_nonempty,
        'envioParsed': envio_parsed,
        'envioSample': envio_sample,
    }


def merge_incoming_rows(incoming, source='upload-xlsx'):
    incoming = [row for row in incoming if (row.get('pedido') or '').strip()]
    with DATA_LOCK:
        current = read_csv_rows(STATE_CSV)
        by_id = {(row.get('pedido') or '').strip(): row for row in current if (row.get('pedido') or '').strip()}
        for row in incoming:
            clean = {key: row.get(key, '') or '' for key in FIELDNAMES}
            clean['pedido'] = (clean['pedido'] or '').strip()
            by_id[clean['pedido']] = clean
        merged = list(by_id.values())
        write_csv_rows(STATE_CSV, merged)
        meta = write_meta(len(merged), source=source)
    return meta, len(incoming)


def choose_data_dir() -> Path:
    configured = os.environ.get('DATA_DIR')
    railway_volume = os.environ.get('RAILWAY_VOLUME_MOUNT_PATH')
    candidates = []
    if configured:
        candidates.append(Path(configured))
    if railway_volume:
        candidates.append(Path(railway_volume))
    if os.environ.get('RAILWAY_ENVIRONMENT') or os.environ.get('RAILWAY_PROJECT_ID'):
        candidates.append(Path('/data'))
    candidates.append(ROOT / 'data' / 'runtime')

    for candidate in candidates:
        try:
            candidate.mkdir(parents=True, exist_ok=True)
            probe = candidate / '.write_test'
            probe.write_text('ok', encoding='utf-8')
            probe.unlink(missing_ok=True)
            return candidate
        except OSError:
            continue
    raise RuntimeError('Nenhum diretório gravável disponível para os dados compartilhados.')


DATA_DIR = choose_data_dir()
STATE_CSV = DATA_DIR / 'shared_dataset.csv'
STATE_META = DATA_DIR / 'shared_meta.json'
DATA_LOCK = threading.RLock()


def utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat().replace('+00:00', 'Z')


def atomic_write_bytes(path: Path, data: bytes) -> None:
    tmp = path.with_suffix(path.suffix + '.tmp')
    tmp.write_bytes(data)
    os.replace(tmp, path)


def atomic_write_text(path: Path, text: str) -> None:
    atomic_write_bytes(path, text.encode('utf-8'))


def read_csv_rows(path: Path):
    if not path.exists():
        return []
    with path.open('r', encoding='utf-8-sig', newline='') as fh:
        return [dict(row) for row in csv.DictReader(fh)]


def write_csv_rows(path: Path, rows) -> None:
    buf = io.StringIO(newline='')
    writer = csv.DictWriter(buf, fieldnames=FIELDNAMES, extrasaction='ignore', lineterminator='\n')
    writer.writeheader()
    for row in rows:
        writer.writerow({key: row.get(key, '') or '' for key in FIELDNAMES})
    atomic_write_text(path, buf.getvalue())


def read_meta():
    try:
        return json.loads(STATE_META.read_text(encoding='utf-8'))
    except Exception:
        rows = read_csv_rows(STATE_CSV)
        return {'version': 'unknown', 'updatedAt': None, 'count': len(rows)}


def write_meta(count: int, source: str = 'upload'):
    meta = {
        'version': str(time.time_ns()),
        'updatedAt': utc_now_iso(),
        'count': count,
        'source': source,
    }
    atomic_write_text(STATE_META, json.dumps(meta, ensure_ascii=False))
    return meta


def ensure_shared_dataset():
    with DATA_LOCK:
        initial_rows = read_csv_rows(INITIAL_DATA) if INITIAL_DATA.exists() else []

        if STATE_CSV.exists():
            current_rows = read_csv_rows(STATE_CSV)

            # Railway Volumes survive redeploys. Older deployments could have
            # persisted only MG, so a new image containing MG + SPN would still
            # keep serving the old MG-only CSV. Reconcile the bundled snapshot
            # by adding only orders that are missing from the persistent base.
            # Existing/shared records always win, so newer uploads/edits are not
            # overwritten by the bundled snapshot.
            if initial_rows:
                current_by_id = {
                    (row.get('pedido') or '').strip(): row
                    for row in current_rows
                    if (row.get('pedido') or '').strip()
                }
                added = 0
                for row in initial_rows:
                    pedido = (row.get('pedido') or '').strip()
                    if pedido and pedido not in current_by_id:
                        current_by_id[pedido] = row
                        added += 1

                if added:
                    current_rows = list(current_by_id.values())
                    write_csv_rows(STATE_CSV, current_rows)
                    write_meta(len(current_rows), source='startup-reconcile')
                    print(
                        f'Reconciliação inicial: {added} pedidos ausentes adicionados '
                        f'(regionais presentes: {sorted({(r.get("reg") or "").strip().upper() for r in current_rows if (r.get("reg") or "").strip()})}).',
                        flush=True,
                    )
                    return

            if not STATE_META.exists():
                write_meta(len(current_rows), source='recovered')
            return

        if not initial_rows:
            write_csv_rows(STATE_CSV, [])
            write_meta(0, source='empty')
            return

        # Normalize the embedded snapshot into the shared schema on first startup.
        write_csv_rows(STATE_CSV, initial_rows)
        write_meta(len(initial_rows), source='initial')


def merge_incoming_csv(csv_text: str):
    reader = csv.DictReader(io.StringIO(csv_text))
    if not reader.fieldnames or 'pedido' not in reader.fieldnames:
        raise ValueError('CSV sem a coluna normalizada "pedido".')

    incoming = []
    for raw in reader:
        pedido = (raw.get('pedido') or '').strip()
        if not pedido:
            continue
        row = {key: (raw.get(key) or '') for key in FIELDNAMES}
        row['pedido'] = pedido
        incoming.append(row)

    with DATA_LOCK:
        current = read_csv_rows(STATE_CSV)
        by_id = {(row.get('pedido') or '').strip(): row for row in current if (row.get('pedido') or '').strip()}
        for row in incoming:
            by_id[row['pedido']] = row
        merged = list(by_id.values())
        write_csv_rows(STATE_CSV, merged)
        meta = write_meta(len(merged), source='upload')
    return meta, len(incoming)


ensure_shared_dataset()


class Handler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def send_json(self, status, payload):
        data = json.dumps(payload, ensure_ascii=False).encode('utf-8')
        self.send_response(status)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Content-Length', str(len(data)))
        self.send_header('Cache-Control', 'no-store, no-cache, must-revalidate')
        self.end_headers()
        self.wfile.write(data)

    def do_GET(self):
        path = urlparse(self.path).path
        if path == '/health':
            meta = read_meta()
            self.send_json(200, {'status': 'ok', 'sharedData': meta})
            return

        if path == '/api/meta':
            with DATA_LOCK:
                meta = read_meta()
            self.send_json(200, meta)
            return

        if path == '/api/data':
            with DATA_LOCK:
                ensure_shared_dataset()
                data = STATE_CSV.read_bytes()
                meta = read_meta()
            self.send_response(200)
            self.send_header('Content-Type', 'text/csv; charset=utf-8')
            self.send_header('Content-Length', str(len(data)))
            self.send_header('Cache-Control', 'no-store, no-cache, must-revalidate')
            self.send_header('X-Data-Version', str(meta.get('version') or ''))
            self.send_header('X-Updated-At', str(meta.get('updatedAt') or ''))
            self.send_header('X-Data-Count', str(meta.get('count') or 0))
            self.end_headers()
            self.wfile.write(data)
            return

        if path in ('/', '/index.html'):
            data = INDEX.read_bytes()
            self.send_response(200)
            self.send_header('Content-Type', 'text/html; charset=utf-8')
            self.send_header('Content-Length', str(len(data)))
            self.send_header('Cache-Control', 'no-cache')
            self.end_headers()
            self.wfile.write(data)
            return

        return super().do_GET()

    def do_POST(self):
        path = urlparse(self.path).path

        try:
            content_length = int(self.headers.get('Content-Length', '0'))
        except ValueError:
            content_length = 0
        if content_length <= 0:
            self.send_json(400, {'error': 'Upload vazio.'})
            return
        if content_length > MAX_UPLOAD_BYTES:
            self.send_json(413, {'error': f'Upload excede o limite de {MAX_UPLOAD_BYTES // (1024 * 1024)} MB.'})
            return

        if path == '/api/import-xlsx':
            try:
                raw = self.rfile.read(content_length)
                rows, diagnostics = parse_xlsx_rows(raw)
                meta, received = merge_incoming_rows(rows, source='upload-xlsx')
                self.send_json(200, {
                    'ok': True,
                    'received': received,
                    'count': meta['count'],
                    'version': meta['version'],
                    'updatedAt': meta['updatedAt'],
                    'diagnostics': diagnostics,
                })
            except ValueError as exc:
                self.send_json(400, {'error': str(exc)})
            except Exception as exc:
                print(f'Erro ao importar XLSX: {exc}', flush=True)
                self.send_json(500, {'error': f'Falha ao importar XLSX: {exc}'})
            return

        if path != '/api/upload':
            self.send_json(404, {'error': 'Rota não encontrada.'})
            return

        try:
            raw = self.rfile.read(content_length)
            text = raw.decode('utf-8-sig')
            meta, received = merge_incoming_csv(text)
            self.send_json(200, {
                'ok': True,
                'received': received,
                'count': meta['count'],
                'version': meta['version'],
                'updatedAt': meta['updatedAt'],
            })
        except UnicodeDecodeError:
            self.send_json(400, {'error': 'Conteúdo enviado não está em UTF-8.'})
        except ValueError as exc:
            self.send_json(400, {'error': str(exc)})
        except Exception as exc:
            print(f'Erro no upload: {exc}', flush=True)
            self.send_json(500, {'error': 'Falha interna ao salvar a base compartilhada.'})

    def log_message(self, fmt, *args):
        print(f'{self.address_string()} - {fmt % args}', flush=True)


if __name__ == '__main__':
    port = int(os.environ.get('PORT', '8080'))
    server = ThreadingHTTPServer(('0.0.0.0', port), Handler)
    print(f'Torre de Controle C2C rodando na porta {port}', flush=True)
    print(f'Dados compartilhados em: {STATE_CSV}', flush=True)
    server.serve_forever()
